from datetime import datetime
from io import BytesIO

import pandas as pd
from fastapi import APIRouter, File, HTTPException, UploadFile

from app.database import get_connection


router = APIRouter(
    prefix="/api/stock",
    tags=["Stock"]
)


# Excel column -> SQL column
COLUMN_MAPPING = {
    "G/F/P": "GFP",
    "Plant": "Plant",
    "Section": "Section",
    "Storage Location": "StorageLocation",
    "Grade": "Grade",
    "Weight (Kg)": "WeightKg",
    "Bin Or Bag No": "BinBagNo",
    "OrderNo:/Issued Order No: (ForExKiln)": "OrderNoIssuedOrderNo",
    "OS AC": "OSAC",
    "Machine": "Machine",
    "Comment": "Comment",
    "Actual Issues": "ActualIssues",
    "corection": "Corection",
    "Qty(Kg)": "QtyKg",
    "SizeCode": "SizeCode",
    "CTC": "CTC",
    "BD": "BD",
    "Date": "Date",
    "Kiln / Machine": "KilnMachine",
    "Origin": "Origin",
    "ExKilnSize": "ExKilnSize",
    "LOCATION": "Location",
    "Catagary": "Catagary",
    "No of Bags": "No of Bags",
    "Mo 6%": "Mo 6%",
    "NMB": "NMB",
    "NMB NO": "NMB NO",
    "Grade1": "Grade1",
    "CTC1": "CTC1",
    "BD1": "BD1",
    "Mo": "Mo",
    "I2": "I2",
    "Ash": "Ash",
    "Sand": "Sand",
    "Magnetic": "Magnetic",
    "Weight (Kg)1": "WeightKg1",
    "In": "In",
}


# Columns that are numeric in SQL Server
NUMERIC_COLUMNS = [
    "GFP",
    "WeightKg",
    "No of Bags",
    "Mo 6%",
    "NMB NO",
    "CTC1",
    "BD1",
    "Mo",
    "I2",
    "Ash",
    "Sand",
    "Magnetic",
    "WeightKg1",
]


# Columns that are dates in SQL Server
DATE_COLUMNS = [
    "Date",
    "In",
]


# Columns that are text in SQL Server
TEXT_COLUMNS = [
    "Plant",
    "Section",
    "StorageLocation",
    "Grade",
    "BinBagNo",
    "OrderNoIssuedOrderNo",
    "OSAC",
    "Machine",
    "Comment",
    "ActualIssues",
    "Corection",
    "QtyKg",
    "SizeCode",
    "CTC",
    "BD",
    "KilnMachine",
    "Origin",
    "ExKilnSize",
    "Location",
    "Catagary",
    "NMB",
    "Grade1",
]


@router.post("/import")
async def import_stock(file: UploadFile = File(...)):
    """
    Import Stock data from an Excel file into Azure SQL.
    """

    # Check file extension
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=400,
            detail="Only Excel files are supported."
        )

    try:
        # ---------------------------------------------------------
        # 1. Read Excel file
        # ---------------------------------------------------------
        contents = await file.read()

        dataframe = pd.read_excel(
            BytesIO(contents),
            sheet_name="Stock",
            dtype=object

        )

        # ---------------------------------------------------------
        # 2. Check required Excel columns
        # ---------------------------------------------------------
        missing_excel_columns = [
            column
            for column in COLUMN_MAPPING
            if column not in dataframe.columns
        ]

        if missing_excel_columns:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": "The Excel file is missing required columns.",
                    "missing_columns": missing_excel_columns
                }
            )

        # ---------------------------------------------------------
        # 3. Keep only the columns we need
        # ---------------------------------------------------------
        dataframe = dataframe[list(COLUMN_MAPPING.keys())]

        # Rename Excel columns to SQL column names
        dataframe = dataframe.rename(columns=COLUMN_MAPPING)

        # ---------------------------------------------------------
        # 4. Convert numeric columns
        # ---------------------------------------------------------
        for column in NUMERIC_COLUMNS:
            dataframe[column] = pd.to_numeric(
                dataframe[column],
                errors="coerce"
            )

        # ---------------------------------------------------------
        # 5. Convert date columns
        # ---------------------------------------------------------
        for column in DATE_COLUMNS:
            dataframe[column] = pd.to_datetime(
                dataframe[column],
                errors="coerce"
            )

        # ---------------------------------------------------------
        # 6. Convert text columns
        # ---------------------------------------------------------
        for column in TEXT_COLUMNS:
            dataframe[column] = dataframe[column].apply(
                lambda value: (
                    None
                    if pd.isna(value)
                    else str(value).strip()
                )
            )

        # ---------------------------------------------------------
        # 7. Convert numeric NaN values to None
        # ---------------------------------------------------------
        #for column in NUMERIC_COLUMNS:
        #    dataframe[column] = dataframe[column].apply(
        #        lambda value: None if pd.isna(value) else float(value)
        #    )

        # ---------------------------------------------------------
        # 8. Convert dates to Python datetime / None
        # ---------------------------------------------------------
        for column in DATE_COLUMNS:
            dataframe[column] = dataframe[column].apply(
                lambda value: (
                    None
                    if pd.isna(value)
                    else value.to_pydatetime()
                )
            )

        # ---------------------------------------------------------
        # 9. Connect to SQL Server
        # ---------------------------------------------------------
        connection = get_connection()
        cursor = connection.cursor()

        # ---------------------------------------------------------
        # 10. SQL INSERT
        # ---------------------------------------------------------
        sql = """
            INSERT INTO dbo.ZStockSheetHI
            (
                GFP,   
                Plant,
                Section,
                StorageLocation,
                Grade,
                WeightKg,
                BinBagNo,
                OrderNoIssuedOrderNo,
                OSAC,
                Machine,
                Comment,
                ActualIssues,
                Corection,
                QtyKg,
                SizeCode,
                CTC,
                BD,
                Date,
                KilnMachine,
                Origin,
                ExKilnSize,
                Location,
                Catagary,
                [No of Bags],
                [Mo 6%],
                NMB,
                [NMB NO],
                Grade1,
                CTC1,
                BD1,
                Mo,
                I2,
                Ash,
                Sand,
                Magnetic,
                WeightKg1,
                [In]
            )
            VALUES
            (
                CAST(? AS int),                -- GFP
                CAST(? AS NVARCHAR(50)),       -- Plant
                CAST(? AS NVARCHAR(100)),      -- Section
                CAST(? AS NVARCHAR(50)),       -- StorageLocation
                CAST(? AS NVARCHAR(100)),      -- Grade
                CAST(? AS FLOAT),              -- WeightKg
                CAST(? AS NVARCHAR(100)),      -- BinBagNo
                CAST(? AS NVARCHAR(100)),      -- OrderNoIssuedOrderNo
                CAST(? AS NVARCHAR(100)),      -- OSAC
                CAST(? AS NVARCHAR(100)),      -- Machine
                CAST(? AS NVARCHAR(500)),      -- Comment
                CAST(? AS NVARCHAR(100)),      -- ActualIssues
                CAST(? AS NVARCHAR(100)),      -- Corection
                CAST(? AS NVARCHAR(100)),      -- QtyKg
                CAST(? AS NVARCHAR(100)),      -- SizeCode
                CAST(? AS NVARCHAR(100)),      -- CTC
                CAST(? AS NVARCHAR(100)),      -- BD
                CAST(? AS DATETIME),           -- Date
                CAST(? AS NVARCHAR(100)),      -- KilnMachine
                CAST(? AS NVARCHAR(100)),      -- Origin
                CAST(? AS NVARCHAR(100)),      -- ExKilnSize
                CAST(? AS NVARCHAR(100)),      -- Location
                CAST(? AS NVARCHAR(100)),      -- Catagary
                CAST(? AS FLOAT),              -- [No of Bags]
                CAST(? AS FLOAT),              -- [Mo 6%]
                CAST(? AS NVARCHAR(100)),      -- NMB
                CAST(? AS FLOAT),              -- [NMB NO]
                CAST(? AS NVARCHAR(100)),      -- Grade1
                CAST(? AS FLOAT),              -- CTC1
                CAST(? AS FLOAT),              -- BD1
                CAST(? AS FLOAT),              -- Mo
                CAST(? AS FLOAT),              -- I2
                CAST(? AS FLOAT),              -- Ash
                CAST(? AS FLOAT),              -- Sand
                CAST(? AS FLOAT),              -- Magnetic
                CAST(? AS FLOAT),              -- WeightKg1
                CAST(? AS DATETIME)            -- In
            )
        """

        # ---------------------------------------------------------
        # 11. Insert rows
        # ---------------------------------------------------------
        inserted_rows = 0

        def clean_value(value):
            """
            Convert pandas/numpy values into normal Python values
            that pyodbc can safely send to SQL Server.
            """

            # Blank / NaN / NaT -> SQL NULL
            if pd.isna(value):
                return None

            # pandas Timestamp -> Python datetime
            if isinstance(value, pd.Timestamp):
                return value.to_pydatetime()

            # numpy numeric types -> normal Python number
            if hasattr(value, "item"):
                try:
                    return value.item()
                except Exception:
                    pass

            return value

        # ---------------------------------------------------------
        # Insert each Excel row
        # ---------------------------------------------------------
        for _, row in dataframe.iterrows():

            # Ignore completely empty rows
            if all(pd.isna(value)for value in row):
                continue

            values = [
                clean_value(row["GFP"]),
                clean_value(row["Plant"]),
                clean_value(row["Section"]),
                clean_value(row["StorageLocation"]),
                clean_value(row["Grade"]),
                clean_value(row["WeightKg"]),
                clean_value(row["BinBagNo"]),
                clean_value(row["OrderNoIssuedOrderNo"]),
                clean_value(row["OSAC"]),
                clean_value(row["Machine"]),
                clean_value(row["Comment"]),
                clean_value(row["ActualIssues"]),
                clean_value(row["Corection"]),
                clean_value(row["QtyKg"]),
                clean_value(row["SizeCode"]),
                clean_value(row["CTC"]),
                clean_value(row["BD"]),
                clean_value(row["Date"]),
                clean_value(row["KilnMachine"]),
                clean_value(row["Origin"]),
                clean_value(row["ExKilnSize"]),
                clean_value(row["Location"]),
                clean_value(row["Catagary"]),
                clean_value(row["No of Bags"]),
                clean_value(row["Mo 6%"]),
                clean_value(row["NMB"]),
                clean_value(row["NMB NO"]),
                clean_value(row["Grade1"]),
                clean_value(row["CTC1"]),
                clean_value(row["BD1"]),
                clean_value(row["Mo"]),
                clean_value(row["I2"]),
                clean_value(row["Ash"]),
                clean_value(row["Sand"]),
                clean_value(row["Magnetic"]),
                clean_value(row["WeightKg1"]),
                clean_value(row["In"]),
            ]

            cursor.execute(sql, values)

            inserted_rows += 1

        # ---------------------------------------------------------
        # 12. Commit transaction
        # ---------------------------------------------------------
        connection.commit()

        cursor.close()
        connection.close()

        return {
            "message": "Excel file imported successfully.",
            "rowsInserted": inserted_rows
        }

    except HTTPException:
        raise

    except Exception as exception:

        # Rollback if something goes wrong
        try:
            connection.rollback()
            cursor.close()
            connection.close()
        except Exception:
            pass

        raise HTTPException(
            status_code=500,
            detail=str(exception)
        )


@router.get("/export")
def export_stock():
    """
    Export ZStockSheetHI table to Excel.
    """

    try:
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                GFP,
                Plant,
                Section,
                StorageLocation,
                Grade,
                WeightKg,
                BinBagNo,
                OrderNoIssuedOrderNo,
                OSAC,
                Machine,
                Comment,
                ActualIssues,
                Corection,
                QtyKg,
                SizeCode,
                CTC,
                BD,
                Date,
                KilnMachine,
                Origin,
                ExKilnSize,
                Location,
                Catagary,
                [No of Bags],
                [Mo 6%],
                NMB,
                [NMB NO],
                Grade1,
                CTC1,
                BD1,
                Mo,
                I2,
                Ash,
                Sand,
                Magnetic,
                WeightKg1,
                [In]
            FROM dbo.ZStockSheetHI
            ORDER BY Date
            """
        )

        rows = cursor.fetchall()

        cursor.close()
        connection.close()

        # Create DataFrame
        columns = [
            "G/F/P",    
            "Plant",
            "Section",
            "Storage Location",
            "Grade",
            "Weight (Kg)",
            "Bin Or Bag No",
            "OrderNo:/Issued Order No: (ForExKiln)",
            "OS AC",
            "Machine",
            "Comment",
            "Actual Issues",
            "corection",
            "Qty(Kg)",
            "SizeCode",
            "CTC",
            "BD",
            "Date",
            "Kiln / Machine",
            "Origin",
            "ExKilnSize",
            "LOCATION",
            "Catagary",
            "No of Bags",
            "Mo 6%",
            "NMB",
            "NMB NO",
            "Grade1",
            "CTC1",
            "BD1",
            "Mo",
            "I2",
            "Ash",
            "Sand",
            "Magnetic",
            "Weight (Kg)1",
            "In",
        ]

        dataframe = pd.DataFrame(
            [list(row) for row in rows],
            columns=columns
        )

        # Create Excel file
        excel_stream = BytesIO()

        with pd.ExcelWriter(
            excel_stream,
            engine="openpyxl"
        ) as writer:

            dataframe.to_excel(
                writer,
                sheet_name="Stock",
                index=False
            )

            worksheet = writer.sheets["Stock"]

            # Freeze header row
            worksheet.freeze_panes = "A2"

            # Add Excel table
            from openpyxl.worksheet.table import Table, TableStyleInfo
            from openpyxl.utils import get_column_letter

            if worksheet.max_row > 1:

                table_reference = (
                    f"A1:"
                    f"{get_column_letter(worksheet.max_column)}"
                    f"{worksheet.max_row}"
                )

                table = Table(
                    displayName="StockTable",
                    ref=table_reference
                )

                table_style = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False
                )

                table.tableStyleInfo = table_style

                worksheet.add_table(table)

            # Auto-size columns
            for column_cells in worksheet.columns:

                maximum_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    if cell.value is not None:
                        maximum_length = max(
                            maximum_length,
                            len(str(cell.value))
                        )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    maximum_length + 2,
                    35
                )

        excel_stream.seek(0)

        filename = (
            f"Stock_Export_"
            f"{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )

        from fastapi.responses import StreamingResponse

        return StreamingResponse(
            excel_stream,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )

    except Exception as exception:

        raise HTTPException(
            status_code=500,
            detail=str(exception)
        )
        
@router.get("")
def get_stock():
    """
    Get all stock data from ZStockSheetHI.
    Used by the React UI to display stock data.
    """

    try:
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                GFP,
                Plant,
                Section,
                StorageLocation,
                Grade,
                WeightKg,
                BinBagNo,
                OrderNoIssuedOrderNo,
                OSAC,
                Machine,
                Comment,
                ActualIssues,
                Corection,
                QtyKg,
                SizeCode,
                CTC,
                BD,
                Date,
                KilnMachine,
                Origin,
                ExKilnSize,
                Location,
                Catagary,
                [No of Bags],
                [Mo 6%],
                NMB,
                [NMB NO],
                Grade1,
                CTC1,
                BD1,
                Mo,
                I2,
                Ash,
                Sand,
                Magnetic,
                WeightKg1,
                [In]
            FROM dbo.ZStockSheetHI
            ORDER BY Date
            """
        )

        rows = cursor.fetchall()
        columns = [column[0] for column in cursor.description]

        cursor.close()
        connection.close()

        data = []

        for row in rows:
            record = {}

            for index, column in enumerate(columns):
                value = row[index]

                # Convert datetime to a format React can display
                if isinstance(value, datetime):
                    value = value.strftime("%Y-%m-%d")

                record[column] = value

            data.append(record)

        return data

    except Exception as exception:
        raise HTTPException(
            status_code=500,
            detail=str(exception)
        )       